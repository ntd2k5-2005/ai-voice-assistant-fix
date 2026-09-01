from __future__ import annotations

from pathlib import Path
import zipfile

import pytest

from rag.config import ParserConfig
from rag.contracts import DocumentSource
from rag.ingestion.parsers import ParserError, ParserErrorCode, ParserRegistry, default_registry


FIXTURES = Path(__file__).parent / "fixtures"


def test_markdown_preserves_headings_list_table_code_and_links():
    documents = default_registry.parse(DocumentSource(str(FIXTURES / "structured.md")))
    block_types = [document.metadata["block_type"] for document in documents]

    assert "list" in block_types
    assert block_types.count("table") == 1
    assert "code" in block_types
    install = next(document for document in documents if document.metadata.get("heading") == "Cài đặt")
    assert install.metadata["heading_path"] == ["Hướng dẫn", "Cài đặt"]
    assert install.metadata["links"] == ["https://example.test/docs"]
    assert "tiếng Việt" in "\n".join(document.content for document in documents)


def test_html_removes_executable_noise_and_keeps_ordered_structures():
    documents = default_registry.parse(DocumentSource(str(FIXTURES / "structured.html")))
    combined = "\n".join(document.content for document in documents)

    assert "window.evil" not in combined
    assert "Bước một" in combined
    table = next(document for document in documents if document.metadata["block_type"] == "table")
    assert table.metadata["column_names"] == ["Tên", "Vai trò"]
    paragraph = next(document for document in documents if "liên kết" in document.content)
    assert paragraph.metadata["links"] == ["https://example.test/link"]
    assert paragraph.metadata["heading_path"] == ["Khởi động"]


def test_jsonl_preserves_falsey_records_and_rejects_partial_corruption(tmp_path):
    documents = default_registry.parse(DocumentSource(str(FIXTURES / "records.jsonl")))
    assert [document.metadata["line_number"] for document in documents] == [1, 2, 3]
    assert documents[1].content == "0"
    assert documents[2].content == "false"

    malformed = tmp_path / "bad.jsonl"
    malformed.write_text('{"ok": 1}\nnot-json\n', encoding="utf-8")
    with pytest.raises(ParserError) as error:
        default_registry.parse(DocumentSource(str(malformed)))
    assert error.value.code is ParserErrorCode.CORRUPT_CONTENT
    assert error.value.diagnostics["line"] == 2


def test_plain_text_keeps_unicode_indentation_and_line_semantics():
    document = default_registry.parse(DocumentSource(str(FIXTURES / "unicode.txt")))[0]
    assert "Đây là tiếng Việt" in document.content
    assert "  nội dung thụt lề" in document.content
    assert "| cột A | cột B |" in document.content
    assert document.metadata["line_start"] == 1
    assert document.metadata["encoding"].lower().replace("_", "-") in {"utf-8", "utf-8-sig"}


def test_docx_preserves_body_order_heading_table_and_merge_provenance(tmp_path):
    from docx import Document

    document = Document()
    document.add_heading("Nhân sự", level=1)
    document.add_paragraph("Đoạn mở đầu tiếng Việt.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Tên"
    table.cell(0, 1).text = "Vai trò"
    table.cell(1, 0).text = "Lan"
    table.cell(1, 1).text = "Kỹ sư"
    table.cell(1, 0).merge(table.cell(1, 1))
    document.add_paragraph("Đoạn sau bảng.")
    path = tmp_path / "ordered.docx"
    document.save(path)

    parsed = default_registry.parse(DocumentSource(str(path)))
    types = [item.metadata["block_type"] for item in parsed]
    assert types == ["section", "table", "paragraph"]
    assert parsed[0].content == "Nhân sự\nĐoạn mở đầu tiếng Việt."
    table_item = parsed[1]
    assert table_item.metadata["heading_path"] == ["Nhân sự"]
    assert table_item.metadata["table_index"] == 1
    assert table_item.metadata["merged_cells"]
    assert parsed[0].metadata["segment_index"] < table_item.metadata["segment_index"] < parsed[2].metadata["segment_index"]



def test_docx_detects_normal_style_bold_section_labels_without_promoting_long_sentences(tmp_path):
    from docx import Document

    document = Document()
    title = document.add_paragraph("Mục tiêu của đề tài:")
    for run in title.runs:
        run.bold = True
    document.add_paragraph(
        "Mục tiêu của nghiên cứu này là xây dựng một hệ thống truy xuất tri thức tiếng Việt "
        "và đoạn văn dài này không được phép bị hiểu nhầm thành heading chỉ vì bắt đầu bằng từ Mục."
    )
    result = document.add_paragraph("Kết quả đạt được:")
    for run in result.runs:
        run.bold = True
    document.add_paragraph("Hệ thống đã hoàn thành hai luồng ingestion và retrieval.")
    path = tmp_path / "normal-style-sections.docx"
    document.save(path)

    parsed = default_registry.parse(DocumentSource(str(path)))
    headings = [item.metadata.get("heading") for item in parsed]
    assert "Mục tiêu của đề tài:" in headings
    assert "Kết quả đạt được:" in headings
    assert not any(
        str(heading).startswith("Mục tiêu của nghiên cứu này là xây dựng")
        for heading in headings
    )
    first = next(item for item in parsed if item.metadata.get("heading") == "Mục tiêu của đề tài:")
    assert "Mục tiêu của nghiên cứu này" in first.content


def test_xlsx_preserves_sheets_sparse_rows_merges_and_formulas(tmp_path):
    from openpyxl import Workbook

    workbook = Workbook()
    employees = workbook.active
    employees.title = "Employees"
    employees.append(["name", "role", "calculation"])
    employees.append(["Lan", "Engineer", "=1+1"])
    employees.merge_cells("A4:B4")
    employees["A4"] = "Summary"
    offices = workbook.create_sheet("Offices")
    offices.append(["city", "country"])
    offices.append(["Hà Nội", "Việt Nam"])
    path = tmp_path / "workbook.xlsx"
    workbook.save(path)

    parsed = default_registry.parse(DocumentSource(str(path)))
    assert {item.metadata["sheet_name"] for item in parsed} == {"Employees", "Offices"}
    formula_row = next(item for item in parsed if item.metadata.get("row_number") == 2 and item.metadata["sheet_name"] == "Employees")
    assert formula_row.metadata["formulas"]["calculation"]["formula"] == "=1+1"
    summary = next(item for item in parsed if item.metadata.get("row_number") == 4)
    assert "A4:B4" in summary.metadata["merged_ranges"]
    assert summary.metadata["row_start"] == summary.metadata["row_end"] == 4


def test_pptx_preserves_slides_tables_notes_and_visual_order(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Kết quả"
    slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(4), Inches(1)).text = "Doanh thu tăng"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(5), Inches(1.5)).table
    table.cell(0, 0).text = "Quý"
    table.cell(0, 1).text = "Giá trị"
    table.cell(1, 0).text = "Q1"
    table.cell(1, 1).text = "100"
    slide.notes_slide.notes_text_frame.text = "Ghi chú diễn giả"
    path = tmp_path / "deck.pptx"
    presentation.save(path)

    parsed = default_registry.parse(DocumentSource(str(path)))
    assert all(item.metadata["slide_number"] == 1 for item in parsed)
    assert any(item.metadata["block_type"] == "table" for item in parsed)
    assert any(item.metadata["block_type"] == "notes" and "Ghi chú" in item.content for item in parsed)
    assert [item.metadata["block_index"] for item in parsed] == sorted(item.metadata["block_index"] for item in parsed)


def test_epub_follows_spine_order_and_keeps_item_provenance(tmp_path):
    path = tmp_path / "book.epub"
    container = """<?xml version='1.0'?>
    <container xmlns='urn:oasis:names:tc:opendocument:xmlns:container'>
      <rootfiles><rootfile full-path='OEBPS/content.opf'/></rootfiles>
    </container>"""
    package = """<?xml version='1.0'?>
    <package xmlns='http://www.idpf.org/2007/opf'>
      <manifest>
        <item id='second' href='second.xhtml' media-type='application/xhtml+xml'/>
        <item id='first' href='first.xhtml' media-type='application/xhtml+xml'/>
      </manifest>
      <spine><itemref idref='first'/><itemref idref='second'/></spine>
    </package>"""
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        archive.writestr("META-INF/container.xml", container)
        archive.writestr("OEBPS/content.opf", package)
        archive.writestr("OEBPS/first.xhtml", "<html><body><h1>First</h1><p>Alpha</p></body></html>")
        archive.writestr("OEBPS/second.xhtml", "<html><body><h1>Second</h1><p>Beta</p></body></html>")

    parsed = default_registry.parse(DocumentSource(str(path)))
    first_beta = next(index for index, item in enumerate(parsed) if "Beta" in item.content)
    first_alpha = next(index for index, item in enumerate(parsed) if "Alpha" in item.content)
    assert first_alpha < first_beta
    assert {item.metadata["spine_index"] for item in parsed} == {1, 2}
    assert parsed[first_alpha].metadata["epub_item"] == "OEBPS/first.xhtml"


def test_signature_dispatch_and_strict_extension_mismatch(tmp_path):
    from docx import Document

    original = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("Signature dispatch")
    document.save(original)
    disguised = tmp_path / "source.pdf"
    disguised.write_bytes(original.read_bytes())

    permissive = default_registry.parse(DocumentSource(str(disguised)))[0]
    assert permissive.file_type == "docx"
    assert permissive.metadata["declared_file_type"] == "pdf"
    assert permissive.metadata["detected_file_type"] == "docx"
    assert permissive.metadata["extension_mismatch"] is True

    strict = ParserRegistry(ParserConfig(reject_extension_mismatch=True))
    with pytest.raises(ParserError) as error:
        strict.parse(DocumentSource(str(disguised)))
    assert error.value.code is ParserErrorCode.TYPE_MISMATCH


def test_empty_corrupt_unsupported_and_size_failures_are_typed(tmp_path):
    empty = tmp_path / "empty.txt"
    empty.write_bytes(b"")
    with pytest.raises(ParserError) as empty_error:
        default_registry.parse(DocumentSource(str(empty)))
    assert empty_error.value.code is ParserErrorCode.EMPTY_CONTENT

    with pytest.raises(ParserError) as corrupt_error:
        default_registry.parse(DocumentSource(str(FIXTURES / "corrupt.pdf")))
    assert corrupt_error.value.code is ParserErrorCode.CORRUPT_CONTENT

    old_doc = tmp_path / "legacy.doc"
    old_doc.write_bytes(b"not-an-ole-document")
    with pytest.raises(ParserError) as unsupported_error:
        default_registry.parse(DocumentSource(str(old_doc)))
    assert unsupported_error.value.code is ParserErrorCode.UNSUPPORTED_FORMAT

    too_large = tmp_path / "large.txt"
    too_large.write_text("0123456789", encoding="utf-8")
    registry = ParserRegistry(ParserConfig(max_file_bytes=5))
    with pytest.raises(ParserError) as size_error:
        registry.parse(DocumentSource(str(too_large)))
    assert size_error.value.code is ParserErrorCode.INVALID_SOURCE


def test_bytes_source_is_bounded_deterministic_and_does_not_leak_temp_paths():
    payload = (FIXTURES / "structured.md").read_bytes()
    source = DocumentSource(payload, source_type="bytes", file_name="memory.md", metadata={"tenant": "qa"})
    first = ParserRegistry().parse(source)
    second = ParserRegistry().parse(source)

    assert [item.document_id for item in first] == [item.document_id for item in second]
    assert all(item.source_type == "bytes" and item.source == "memory.md" for item in first)
    assert all(item.metadata["source_path"] == "memory.md" for item in first)
    assert all(item.metadata["source_in_memory"] is True for item in first)
    assert all("rag-parser-" not in str(item.metadata) for item in first)
    assert {"md", "jsonl", "epub", "pdf"}.issubset(set(ParserRegistry().supported_types))

    with pytest.raises(ParserError):
        ParserRegistry().parse(DocumentSource(payload, source_type="bytes"))
    with pytest.raises(ParserError):
        ParserRegistry().parse(DocumentSource(payload, source_type="bytes", file_name="../escape.md"))


def test_parser_config_encoding_and_sniffing_propagate_to_runtime(tmp_path):
    encoded = tmp_path / "legacy.txt"
    encoded.write_bytes("café".encode("cp1252"))
    parsed = ParserRegistry(ParserConfig(encoding_fallbacks=("cp1252",))).parse(DocumentSource(str(encoded)))[0]
    assert parsed.content == "café"
    assert parsed.metadata["encoding"] == "cp1252"

    from docx import Document

    disguised = tmp_path / "disguised.pdf"
    document = Document()
    document.add_paragraph("OOXML content")
    document.save(disguised)
    with pytest.raises(ParserError) as error:
        ParserRegistry(ParserConfig(sniff_mime=False)).parse(DocumentSource(str(disguised)))
    assert error.value.code is ParserErrorCode.CORRUPT_CONTENT


def test_nested_json_and_textual_aliases_preserve_content_and_context(tmp_path):
    nested = tmp_path / "nested.json"
    nested.write_text(
        '{"outer":{"items":[{"name":"Hồng ngọc","enabled":false,"count":0}]}}',
        encoding="utf-8",
    )
    parsed = default_registry.parse(DocumentSource(nested))
    assert parsed and parsed[0].metadata["json_path"] == "$.outer"
    assert '"enabled": false' in parsed[0].content
    assert '"count": 0' in parsed[0].content

    aliases = {
        "records.ldjson": '{"id":1}\n{"id":2}\n',
        "page.mdx": "# Tiêu đề\n\nNội dung MDX thuần.",
        "rows.tsv": "name\tvalue\nalpha\t0\n",
    }
    for name, content in aliases.items():
        path = tmp_path / name
        path.write_text(content, encoding="utf-8")
        units = default_registry.parse(DocumentSource(path))
        assert units and any("0" in item.content or "Nội dung" in item.content or '"id"' in item.content for item in units)


@pytest.mark.parametrize(
    "extension",
    ["py", "js", "ts", "java", "c", "cpp", "h", "php", "go", "sh", "cs", "kt", "sql"],
)
def test_code_text_aliases_are_actually_dispatchable(tmp_path, extension):
    path = tmp_path / f"source.{extension}"
    path.write_text("function_marker = 'đã kiểm chứng'\n", encoding="utf-8")
    parsed = default_registry.parse(DocumentSource(path))
    assert parsed and "đã kiểm chứng" in parsed[0].content


def test_xlsm_alias_and_all_legacy_office_extensions(tmp_path):
    from openpyxl import Workbook

    workbook_path = tmp_path / "macro-compatible.xlsm"
    workbook = Workbook()
    workbook.active.append(["name", "value"])
    workbook.active.append(["alpha", 1])
    workbook.save(workbook_path)
    parsed = default_registry.parse(DocumentSource(workbook_path))
    assert parsed and parsed[0].file_type == "xlsx"

    for extension in ("doc", "xls", "ppt"):
        path = tmp_path / f"legacy.{extension}"
        path.write_bytes(b"legacy-binary-placeholder")
        with pytest.raises(ParserError) as error:
            default_registry.parse(DocumentSource(path))
        assert error.value.code is ParserErrorCode.UNSUPPORTED_FORMAT
