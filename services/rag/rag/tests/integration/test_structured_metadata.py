from pathlib import Path

from rag.contracts import DocumentSource
from rag.ingestion.parsers import default_registry
from rag.lifecycle import DocumentRegistry
from rag.pipeline import RetrievalPipeline
from rag.retrieval import SQLiteBM25Retriever
from rag.storage.qdrant import QdrantStore


class DeterministicEmbedder:
    embedding_dim = 2
    model_name = "test-structured-embedder"

    def embed_documents(self, texts):
        return [[1.0, 0.0] for _ in texts]

    def embed_query(self, text):
        return [1.0, 0.0]


def test_format_parsers_preserve_structural_provenance(tmp_path):
    from openpyxl import Workbook
    from pptx import Presentation

    csv_path = tmp_path / "records.csv"
    csv_path.write_text("name,city\nLan,Hanoi\n", encoding="utf-8")
    csv_doc = default_registry.parse(DocumentSource(str(csv_path)))[0]
    assert csv_doc.metadata["row_number"] == 2
    assert csv_doc.metadata["column_names"] == ["name", "city"]

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Employees"
    sheet.append(["name", "role"])
    sheet.append(["Lan", "Engineer"])
    xlsx_path = tmp_path / "records.xlsx"
    workbook.save(xlsx_path)
    xlsx_doc = default_registry.parse(DocumentSource(str(xlsx_path)))[0]
    assert xlsx_doc.metadata["sheet_name"] == "Employees"
    assert xlsx_doc.metadata["row_number"] == 2

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, 100, 100).text = "Quarterly result"
    pptx_path = tmp_path / "deck.pptx"
    presentation.save(pptx_path)
    pptx_doc = default_registry.parse(DocumentSource(str(pptx_path)))[0]
    assert pptx_doc.metadata["slide_number"] == 1

    html_path = tmp_path / "page.html"
    html_path.write_text("<title>Guide</title><h1>Install</h1><p>Use structured metadata.</p>", encoding="utf-8")
    html_doc = default_registry.parse(DocumentSource(str(html_path)))[0]
    assert html_doc.metadata["heading"] == "Install"
    assert html_doc.metadata["source_path"].endswith("page.html")


def test_structured_metadata_survives_to_bm25_result(tmp_path):
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Employees"
    sheet.append(["name", "role"])
    sheet.append(["Lan", "Engineer"])
    source_path = tmp_path / "employees.xlsx"
    workbook.save(source_path)

    store = QdrantStore("structured", 2, str(tmp_path / "qdrant"), embedding_model="test-structured-embedder")
    registry = DocumentRegistry(tmp_path / "registry.db")
    sparse = SQLiteBM25Retriever(tmp_path / "fts.db")
    pipeline = RetrievalPipeline(
        DeterministicEmbedder(),
        store,
        sparse_store=sparse,
        registry=registry,
    )
    try:
        pipeline.ingest_sources([DocumentSource(str(source_path), metadata={"document_id": "employees", "tenant": "qa"})])

        result = pipeline.retrieve_result("Engineer", mode="sparse", top_k=1)
        metadata = result.candidates[0].metadata
        assert metadata["sheet_name"] == "Employees"
        assert metadata["row_number"] == 2
        assert metadata["tenant"] == "qa"
        assert metadata["chunk_index"] == 0
        assert metadata["chunk_id"] == result.candidates[0].chunk_id
    finally:
        pipeline.close()
