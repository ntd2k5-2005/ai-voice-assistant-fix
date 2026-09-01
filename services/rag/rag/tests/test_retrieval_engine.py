from __future__ import annotations

import pytest

from rag.Embeddings.adapters import EngineEmbeddingAdapter
from rag.contracts import DocumentSource, MetadataFilter
from rag.ingestion.parsers import TextParser
from rag.ingestion.parsers.registry import FileParser
from rag.lifecycle import DocumentRegistry
from rag.metrics import aggregate_metrics, mrr
from rag.pipeline import RetrievalPipeline
from rag.retrieval import (
    DenseRetriever,
    HybridRetriever,
    SparseRetriever,
    SQLiteBM25Retriever,
)
from rag.storage.qdrant import QdrantStore


CHUNKS = [
    {
        "_id": "a",
        "document_id": "vn-a",
        "version": 1,
        "text": "Hà Nội là thủ đô Việt Nam",
        "metadata": {"document": "vn", "document_id": "vn-a", "version": 1},
    },
    {
        "_id": "b",
        "document_id": "vn-b",
        "version": 1,
        "text": "Đà Nẵng là thành phố biển",
        "metadata": {"document": "vn", "document_id": "vn-b", "version": 1},
    },
]


class DeterministicEmbedder:
    embedding_dim = 3
    model_name = "tests/deterministic"
    normalize = True

    def __init__(self) -> None:
        self.fail = False
        self.query_calls: list[str] = []

    def embed_documents(self, texts):
        if self.fail:
            raise RuntimeError("embedding unavailable")
        return [[1.0, 0.0, 0.0] for _ in texts]

    def embed_query(self, text):
        self.query_calls.append(text)
        return [1.0, 0.0, 0.0]


def make_pipeline(tmp_path, embedder=None, *, collection="pipeline"):
    embedder = embedder or DeterministicEmbedder()
    store = QdrantStore(
        collection,
        3,
        str(tmp_path / "qdrant"),
        embedding_model=embedder.model_name,
        normalization=True,
    )
    registry = DocumentRegistry(tmp_path / "registry.db")
    sparse = SQLiteBM25Retriever(tmp_path / "bm25.db")
    pipeline = RetrievalPipeline(
        embedder,
        store,
        sparse_store=sparse,
        registry=registry,
    )
    return pipeline, store, registry, sparse


def test_bm25_search_and_metadata_filter():
    result = SparseRetriever(CHUNKS).search(
        "thủ đô", filters=MetadataFilter("document", "=", "vn")
    )
    assert result[0].chunk_id == "a"
    assert result[0].sparse_score is not None


def test_text_parser_preserves_structured_document_contract():
    parsed = TextParser().parse(
        DocumentSource("Xin chào", "text", {"document_id": "raw-1"})
    )
    assert parsed[0].document_id == "raw-1"
    assert parsed[0].content == "Xin chào"


def test_file_parser_formats(tmp_path):
    from docx import Document as DocxDocument
    from openpyxl import Workbook
    from pptx import Presentation
    from reportlab.pdfgen import canvas

    (tmp_path / "a.txt").write_text("Xin chào", encoding="utf-8")
    (tmp_path / "a.csv").write_text("name,value\nfoo,1", encoding="utf-8")
    (tmp_path / "a.html").write_text(
        "<html><script>x</script><h1>Tiêu đề</h1><p>Nội dung</p></html>",
        encoding="utf-8",
    )
    book = Workbook()
    book.active.append(["name", "value"])
    book.active.append(["foo", 1])
    book.save(tmp_path / "a.xlsx")
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, 100, 100).text = "Slide text"
    presentation.save(tmp_path / "a.pptx")
    docx = DocxDocument()
    docx.add_heading("Heading", level=1)
    docx.add_paragraph("Document paragraph content")
    docx.save(tmp_path / "a.docx")
    pdf_path = tmp_path / "a.pdf"
    pdf = canvas.Canvas(str(pdf_path))
    pdf.drawString(
        50, 750, "This is a sufficiently long PDF text fixture for parser validation."
    )
    pdf.save()

    parser = FileParser()
    assert any(
        "Document paragraph" in unit.content
        for unit in parser.parse(DocumentSource(tmp_path / "a.docx"))
    )
    assert any(
        "sufficiently long PDF" in unit.content
        for unit in parser.parse(DocumentSource(pdf_path))
    )
    for name, expected in [
        ("a.txt", "Xin chào"),
        ("a.csv", "foo"),
        ("a.html", "Nội dung"),
        ("a.xlsx", "foo"),
        ("a.pptx", "Slide text"),
    ]:
        result = parser.parse(DocumentSource(tmp_path / name))
        assert any(expected in unit.content for unit in result)


def test_parser_is_connected_to_canonical_ingestion_path(tmp_path):
    pipeline, store, registry, _ = make_pipeline(tmp_path, collection="parser-path")
    try:
        result = pipeline.ingest_sources(
            [
                DocumentSource(
                    "Structured text with enough words for deterministic ingestion.",
                    "text",
                    {"document_id": "raw"},
                )
            ]
        )
        assert result
        assert store.count_document("raw", version=1) == len(result)
        assert registry.get("raw")["status"] == "active"
    finally:
        pipeline.close()


def test_persistent_bm25_incremental_update_and_delete(tmp_path):
    index = SQLiteBM25Retriever(tmp_path / "bm25.db")
    index.upsert(CHUNKS)
    assert index.search("thủ đô")[0].chunk_id == "a"
    index.close()

    index = SQLiteBM25Retriever(tmp_path / "bm25.db")
    assert index.search("thủ đô")[0].chunk_id == "a"
    index.upsert(
        [
            {
                "_id": "a",
                "document_id": "vn-a",
                "version": 1,
                "text": "Nội dung đã cập nhật",
                "metadata": {"document_id": "vn-a", "version": 1},
            }
        ]
    )
    assert not index.search("thủ đô")
    index.delete("b")
    assert not index.search("biển")
    index.close()


def test_registry_versions_and_metrics(tmp_path):
    registry = DocumentRegistry(tmp_path / "registry.db")
    assert registry.register("doc", "a.txt", "v1")["version"] == 1
    assert registry.register("doc", "a.txt", "v2")["version"] == 2
    registry.mark_deleted("doc")
    assert registry.get("doc")["status"] == "deleted"
    assert mrr(["x", "a"], {"a"}) == 0.5
    metrics = aggregate_metrics([{"retrieved": ["x", "a"], "relevant": {"a"}}])
    assert metrics["recall@5"] == 1
    assert metrics["mrr"] == 0.5
    assert aggregate_metrics(
        [{"retrieved": ["a"], "relevant": {"a", "b"}}]
    )["recall@5"] == 0.5
    registry.close()

    registry = DocumentRegistry(tmp_path / "registry.db")
    assert registry.get("doc")["status"] == "deleted"
    registry.close()


def test_embedding_adapter_and_hybrid_reranker_fallback():
    class Engine:
        def embed_documents(self, texts):
            return [[1.0, 0.0] for _ in texts]

        def embed_query(self, text):
            return [1.0, 0.0]

    class Store:
        def search(self, vector, top_k=10, query_filter=None):
            return [
                {"_id": "a", "text": "thủ đô", "score": 0.8, "metadata": {}}
            ]

    class BrokenReranker:
        def rerank_candidates(self, candidates, query, top_k=10):
            raise RuntimeError("model unavailable")

    adapter = EngineEmbeddingAdapter(Engine())
    assert len(adapter.embed_documents(["x"])) == 1
    assert len(adapter.embed_query("x")) == 2
    result = HybridRetriever(
        DenseRetriever(adapter, Store()),
        SparseRetriever(CHUNKS),
        BrokenReranker(),
    ).retrieve("thủ đô", rerank=True)
    assert result.diagnostics["fallback"] is True
    assert result.candidates[0].final_score is not None


def test_pipeline_indexes_both_stores_and_lifecycle(tmp_path):
    pipeline, store, registry, sparse = make_pipeline(tmp_path)
    try:
        first_text = "Nội dung phiên bản một có nhiều từ để kiểm tra vòng đời."
        second_text = "Nội dung phiên bản hai đã thay đổi hoàn toàn cho cập nhật."
        pipeline.ingest(
            [{"document_id": "doc-1", "source": "a.txt", "text": first_text}]
        )
        assert pipeline.retrieve("phiên bản", mode="bm25")[0]["document_id"] == "doc-1"
        dense_result = pipeline.retrieve_result("phiên bản", mode="dense")
        assert dense_result.candidates[0].retrieval_source == "dense"
        assert registry.get("doc-1")["version"] == 1

        pipeline.ingest(
            [{"document_id": "doc-1", "source": "a.txt", "text": first_text}]
        )
        assert pipeline.last_ingestion_results[0].action == "unchanged"
        assert registry.get("doc-1")["version"] == 1

        pipeline.ingest(
            [{"document_id": "doc-1", "source": "a.txt", "text": second_text}]
        )
        assert pipeline.last_ingestion_results[0].action == "update"
        assert registry.get("doc-1")["version"] == 2
        assert store.count_document("doc-1", version=1) == 0
        assert sparse.count_document("doc-1", version=1) == 0
        assert not pipeline.retrieve("một", mode="bm25")

        pipeline.delete_document("doc-1")
        assert registry.get("doc-1")["status"] == "deleted"
        assert not pipeline.retrieve("hai", mode="bm25")
    finally:
        pipeline.close()


def test_ingestion_failure_marks_registry_failed(tmp_path):
    embedder = DeterministicEmbedder()
    embedder.fail = True
    pipeline, _, registry, _ = make_pipeline(
        tmp_path, embedder, collection="broken-ingestion"
    )
    try:
        with pytest.raises(RuntimeError, match="embedding unavailable"):
            pipeline.ingest(
                [
                    {
                        "document_id": "broken",
                        "source": "x.txt",
                        "text": "Content with enough words to reach the embedding stage.",
                    }
                ]
            )
        assert registry.get("broken")["status"] == "failed"
    finally:
        pipeline.close()


def test_update_embedding_failure_preserves_active_version(tmp_path):
    embedder = DeterministicEmbedder()
    pipeline, store, registry, sparse = make_pipeline(
        tmp_path, embedder, collection="stable-update"
    )
    try:
        pipeline.ingest(
            [
                {
                    "document_id": "stable",
                    "source": "x.txt",
                    "text": "Old content remains visible while a later update fails.",
                }
            ]
        )
        embedder.fail = True
        with pytest.raises(RuntimeError, match="embedding unavailable"):
            pipeline.ingest(
                [
                    {
                        "document_id": "stable",
                        "source": "x.txt",
                        "text": "New content cannot replace the active version after failure.",
                    }
                ]
            )
        row = registry.get("stable")
        assert row["version"] == 1
        assert row["status"] == "active"
        assert store.count_document("stable", version=1) > 0
        assert sparse.count_document("stable", version=1) > 0
    finally:
        pipeline.close()
