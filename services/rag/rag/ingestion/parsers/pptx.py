from __future__ import annotations

from ...config import ParserConfig
from .base import (
    DocumentParser, ParserError, ParserErrorCode, parsed_documents, source_path,
    validate_zip_archive,
)


class PPTXParser(DocumentParser):
    """Parse slide shapes in visual reading order with table/notes identity."""

    def __init__(self, config: ParserConfig | None = None):
        self.config = config or ParserConfig()

    def parse(self, source):
        path = source_path(source)
        validate_zip_archive(path, self.config)
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        try:
            presentation = Presentation(str(path))
        except Exception as exc:
            raise ParserError(
                ParserErrorCode.CORRUPT_CONTENT,
                f"cannot open presentation: {type(exc).__name__}: {exc}",
                source=path,
                stage="open",
            ) from exc

        segments: list[dict] = []

        def sorted_shapes(shapes):
            return sorted(
                shapes,
                key=lambda shape: (
                    int(getattr(shape, "top", 0) or 0),
                    int(getattr(shape, "left", 0) or 0),
                ),
            )

        def shape_box(shape) -> dict[str, int]:
            return {
                "left": int(getattr(shape, "left", 0) or 0),
                "top": int(getattr(shape, "top", 0) or 0),
                "width": int(getattr(shape, "width", 0) or 0),
                "height": int(getattr(shape, "height", 0) or 0),
            }

        for slide_number, slide in enumerate(presentation.slides, 1):
            title_shape = slide.shapes.title
            title = (
                title_shape.text.strip()
                if title_shape is not None and getattr(title_shape, "text", "").strip()
                else f"Slide {slide_number}"
            )
            table_index = 0
            block_index = 0

            def add(text: str, block_type: str, shape=None, **extra):
                nonlocal block_index
                text = text.strip()
                if not text:
                    return
                block_index += 1
                segments.append(
                    {
                        "text": text,
                        "title": title,
                        "method": "pptx",
                        "slide_number": slide_number,
                        "page": slide_number,
                        "page_number": slide_number,
                        "heading": title,
                        "heading_path": [title],
                        "section": f"slide {slide_number}",
                        "block_type": block_type,
                        "block_index": block_index,
                        "bounding_box_emu": shape_box(shape) if shape is not None else None,
                        **extra,
                    }
                )

            def visit(shape):
                nonlocal table_index
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    for child in sorted_shapes(shape.shapes):
                        visit(child)
                    return
                if getattr(shape, "has_table", False):
                    table_index += 1
                    rows = [[cell.text.strip().replace("\n", " ") for cell in row.cells] for row in shape.table.rows]
                    if rows:
                        width = max(len(row) for row in rows)
                        rows = [row + [""] * (width - len(row)) for row in rows]
                        header = rows[0]
                        lines = [
                            "| " + " | ".join(header) + " |",
                            "| " + " | ".join(["---"] * width) + " |",
                        ]
                        lines.extend("| " + " | ".join(row) + " |" for row in rows[1:])
                        add(
                            "\n".join(lines),
                            "table",
                            shape,
                            method="pptx-table",
                            table_id=f"slide-{slide_number}-table-{table_index}",
                            table_index=table_index,
                            column_names=header,
                            row_start=1,
                            row_end=len(rows),
                        )
                    return
                if getattr(shape, "has_text_frame", False):
                    paragraphs = []
                    for paragraph in shape.text_frame.paragraphs:
                        value = paragraph.text.strip()
                        if not value:
                            continue
                        try:
                            bulleted = bool(paragraph._p.xpath("./a:pPr/a:buChar | ./a:pPr/a:buAutoNum | ./a:pPr/a:buBlip"))
                        except Exception:
                            bulleted = False
                        if bulleted:
                            value = f"{'  ' * paragraph.level}- {value}"
                        paragraphs.append(value)
                    value = "\n".join(paragraphs)
                    add(value, "title" if shape is title_shape else "text", shape)
                    return
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    description = ""
                    try:
                        values = shape._element.xpath(".//*[local-name()='cNvPr']/@descr | .//*[local-name()='cNvPr']/@title")
                        description = next((str(value).strip() for value in values if str(value).strip()), "")
                    except Exception:
                        pass
                    if description:
                        add(description, "image_caption", shape)

            for shape in sorted_shapes(slide.shapes):
                visit(shape)

            try:
                notes = slide.notes_slide.notes_text_frame.text.strip()
            except (AttributeError, ValueError):
                notes = ""
            note_lines = [
                line.strip() for line in notes.splitlines()
                if line.strip() and line.strip() != str(slide_number)
            ]
            if note_lines:
                add("\n".join(note_lines), "notes")

        if not segments:
            raise ParserError(ParserErrorCode.EMPTY_CONTENT, "presentation contains no readable content", source=path)
        return parsed_documents(source, path, segments)
