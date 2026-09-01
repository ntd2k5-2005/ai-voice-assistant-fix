from __future__ import annotations

from datetime import date, datetime, time

from ...config import ParserConfig
from .base import (
    DocumentParser, ParserError, ParserErrorCode, parsed_documents, source_path,
    validate_zip_archive,
)


def _display(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value).strip()


def _headers(values: list[str], columns: list[int]) -> list[str]:
    result: list[str] = []
    seen: dict[str, int] = {}
    for value, column in zip(values, columns):
        base = value.strip() or f"column_{column}"
        seen[base] = seen.get(base, 0) + 1
        result.append(base if seen[base] == 1 else f"{base}_{seen[base]}")
    return result


class XLSXParser(DocumentParser):
    """Parse sparse workbooks as row units, preserving sheets and formulas."""

    def __init__(self, config: ParserConfig | None = None):
        self.config = config or ParserConfig()

    def parse(self, source):
        path = source_path(source)
        validate_zip_archive(path, self.config)
        from openpyxl import load_workbook

        try:
            formula_book = load_workbook(path, read_only=False, data_only=False, keep_links=False)
            value_book = load_workbook(path, read_only=False, data_only=True, keep_links=False)
        except Exception as exc:
            raise ParserError(
                ParserErrorCode.CORRUPT_CONTENT,
                f"cannot open workbook: {type(exc).__name__}: {exc}",
                source=path,
                stage="open",
            ) from exc

        try:
            segments: list[dict] = []
            for sheet_index, sheet in enumerate(formula_book.worksheets, 1):
                value_sheet = value_book[sheet.title]
                populated = {
                    (row, column): cell
                    for (row, column), cell in sheet._cells.items()
                    if cell.value is not None and _display(cell.value)
                }
                if not populated:
                    continue
                row_numbers = sorted({row for row, _ in populated})
                column_numbers = sorted({column for _, column in populated})
                header_number = row_numbers[0]
                merged_ranges = [str(cell_range) for cell_range in sheet.merged_cells.ranges]
                merge_objects = list(sheet.merged_cells.ranges)

                def effective_value(row: int, column: int):
                    value = sheet.cell(row, column).value
                    if value is not None:
                        return value
                    for merged in merge_objects:
                        if merged.min_row <= row <= merged.max_row and merged.min_col <= column <= merged.max_col:
                            return sheet.cell(merged.min_row, merged.min_col).value
                    return None

                raw_headers = [_display(effective_value(header_number, column)) for column in column_numbers]
                headers = _headers(raw_headers, column_numbers)
                common = {
                    "method": "xlsx",
                    "sheet_name": sheet.title,
                    "sheet_index": sheet_index,
                    "sheet_state": sheet.sheet_state,
                    "section": sheet.title,
                    "table_id": f"sheet-{sheet_index}",
                    "column_names": headers,
                    "column_numbers": column_numbers,
                    "header_row": header_number,
                    "merged_ranges": merged_ranges,
                }

                data_rows = [row for row in row_numbers if row != header_number]
                if not data_rows:
                    segments.append(
                        {
                            **common,
                            "text": "\n".join(headers),
                            "row_number": header_number,
                            "row_start": header_number,
                            "row_end": header_number,
                            "block_type": "table_header",
                        }
                    )
                    continue

                for row_number in data_rows:
                    values: list[str] = []
                    formula_metadata: dict[str, dict[str, str | None]] = {}
                    for header, column in zip(headers, column_numbers):
                        formula_value = effective_value(row_number, column)
                        cached_value = value_sheet.cell(row_number, column).value
                        rendered = _display(formula_value)
                        if isinstance(formula_value, str) and formula_value.startswith("="):
                            formula_metadata[header] = {
                                "formula": formula_value,
                                "cached_value": _display(cached_value) or None,
                            }
                            if cached_value is not None:
                                rendered = f"{formula_value} (value: {_display(cached_value)})"
                        values.append(rendered)
                    if not any(values):
                        continue
                    segments.append(
                        {
                            **common,
                            "text": "\n".join(f"{header}: {value}" for header, value in zip(headers, values)),
                            "row_number": row_number,
                            "row_start": row_number,
                            "row_end": row_number,
                            "block_type": "table_row",
                            "formulas": formula_metadata,
                        }
                    )
            if not segments:
                raise ParserError(ParserErrorCode.EMPTY_CONTENT, "workbook contains no usable cells", source=path)
            return parsed_documents(source, path, segments)
        finally:
            formula_book.close()
            value_book.close()
